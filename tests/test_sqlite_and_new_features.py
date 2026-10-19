"""Tests for SQLite migration, Visa settlement follow-up, per-group instructor rates, and granular user permissions."""
from datetime import date, time, timedelta
from openpyxl import load_workbook
import pytest

import config
from core.audit import AuditService
from core.auth import AuthError, AuthService, Session
from core.excel_repo import ExcelRepository
from core.schemas import Role, Sheet, Student
from services.common import ServiceError
from services.registry import Services

TODAY = date.today()


def test_sqlite_crud_migration_from_excel_and_snapshot_export(tmp_path):
    # 1. Create an initial Excel workbook with a student
    xlsx_path = tmp_path / "academy.xlsx"
    xlsx_repo = ExcelRepository(xlsx_path, tmp_path / "backups_xlsx", current_user=lambda: "admin")
    xlsx_repo.ensure_workbook()
    xlsx_repo.append(Sheet.STUDENTS, Student(student_id=1, name_ar="مريم طارق علي", parent_mobile="01012345678"))

    # 2. Initialize SQLite repository in the same directory -> should auto-migrate from academy.xlsx
    db_path = tmp_path / "academy.db"
    db_repo = ExcelRepository(db_path, tmp_path / "backups_db", current_user=lambda: "admin")
    db_repo.ensure_workbook()

    students = db_repo.read(Sheet.STUDENTS)
    assert len(students) == 1 and students[0].name_ar == "مريم طارق علي"

    # 3. Transaction + backup + rollback on SQLite
    with db_repo.transaction():
        db_repo.append(Sheet.STUDENTS, Student(student_id=2, name_ar="يوسف سامي حسن", parent_mobile="01122334455"))
    assert len(db_repo.read(Sheet.STUDENTS)) == 2
    assert len(db_repo.list_backups()) >= 1

    # 4. Export full Excel snapshot from SQLite
    snap = db_repo.export_excel_snapshot(tmp_path / "snapshot.xlsx")
    wb = load_workbook(snap)
    assert wb.sheetnames == [s.value for s in Sheet]
    assert wb["Students"].max_row == 3  # header + 2 students


def test_visa_payment_settlement_and_per_group_instructor_rate(tmp_path):
    db_path = tmp_path / "academy.db"
    repo = ExcelRepository(db_path, tmp_path / "backups", current_user=lambda: "owner")
    repo.ensure_workbook()
    AuditService(repo, lambda: "owner")
    svc = Services(repo, lambda: "owner")

    # Instructor with base rate 100/h, assigned to two groups with different hourly rates (140/h and 180/h)
    iid = svc.instructors.add("م. أحمد فؤاد", "01099887766", 100)
    g1 = svc.groups.add("بايثون", "برمجة", "مبتدئ", iid, "السبت 5م", "onsite", instructor_hourly_rate=140)
    g2 = svc.groups.add("ذكاء اصطناعي", "برمجة", "متقدم", iid, "الأحد 6م", "online", instructor_hourly_rate=180)

    # Create 1-hour session in each group
    s1 = svc.attendance.create_session(g1, TODAY, time(16, 0), time(17, 0))
    s2 = svc.attendance.create_session(g2, TODAY, time(18, 0), time(19, 0))

    sess1 = repo.get(Sheet.SESSIONS, session_id=s1)
    sess2 = repo.get(Sheet.SESSIONS, session_id=s2)
    assert sess1.rate_applied == 140 and sess1.amount == 140
    assert sess2.rate_applied == 180 and sess2.amount == 180

    summary = svc.instructors.summary(TODAY.replace(day=1), TODAY)[0]
    assert summary["gross"] == 320 and "140" in summary["groups_info"] and "180" in summary["groups_info"]

    # Register student & pay via Visa
    sid = svc.students.register(
        dict(name_ar="سيف الدين خالد", parent_mobile="01234567890"),
        enrollment=dict(group_id=g1, start_date=TODAY, monthly_fee=1200, months=2),
    )
    eid = svc.ledger().student_enrollments(sid)[0].enrollment_id

    rec = svc.receipts.create(
        student_id=sid, enrollment_id=eid, amount=1200, method="visa",
        cost_center="بايثون", description="اشتراك شهري", bank_ref="POS-7788",
    )
    assert rec.method == "visa" and rec.settlement_status == "pending" and rec.bank_ref == "POS-7788"

    sm = svc.receipts.settlement_summary(method="visa")
    assert sm["pending_amount"] == 1200 and sm["settled_amount"] == 0 and sm["pending_count"] == 1

    # Confirm Visa settlement into academy bank account
    updated = svc.receipts.set_settlement(rec.serial, "settled", bank_ref="BANK-OK-99")
    assert updated.settlement_status == "settled" and updated.bank_ref == "BANK-OK-99" and updated.settled_by == "owner"

    sm_after = svc.receipts.settlement_summary(method="visa")
    assert sm_after["pending_amount"] == 0 and sm_after["settled_amount"] == 1200

    # Attendance check_in, check_out, notes & Expense recipient/attachment
    svc.attendance.mark(s1, sid, "present", notes="ممتاز", check_in=time(16, 2), check_out=time(17, 0))
    att = repo.get(Sheet.ATTENDANCE, session_id=s1, student_id=sid)
    assert att.check_in == time(16, 2) and att.check_out == time(17, 0) and att.notes == "ممتاز"

    xid = svc.journal.add_expense(350, "فاتورة إنترنت", "مرافق (كهرباء/مياه/إنترنت)", recipient="WE", attachment="INV-404")
    exp = repo.get(Sheet.EXPENSES, id=xid)
    assert exp.recipient == "WE" and exp.attachment == "INV-404"


def test_granular_user_roles_and_permissions(tmp_path):
    session = Session()
    repo = ExcelRepository(tmp_path / "academy.db", tmp_path / "backups", current_user=lambda: session.username)
    repo.ensure_workbook()
    audit = AuditService(repo, lambda: session.username)
    auth = AuthService(repo, audit, session)

    auth.first_run_setup(("owner", "Owner#12345"), ("reception", "Recep#12345"))
    auth.login("owner", "Owner#12345")

    # Create an accountant user with custom permissions (including reports & settlements)
    auth.create_user("acc_user", "Acc#123456", Role.ACCOUNTANT, permissions=["receipts", "settlements", "journal", "reports"])
    auth.logout()

    auth.login("acc_user", "Acc#123456")
    assert session.can("settlements") and session.can("reports") and session.can("receipts")
    assert not session.can("users") and not session.can("instructors")

    # Self-service password change for any account (accountant, reception, owner)
    with pytest.raises(AuthError, match="غير صحيحة"):
        auth.change_password("acc_user", "WrongOld#1", "NewAcc#98765")
    with pytest.raises(AuthError, match="تختلف"):
        auth.change_password("acc_user", "Acc#123456", "Acc#123456")

    auth.change_password("acc_user", "Acc#123456", "NewAcc#98765")
    auth.logout()
    auth.login("acc_user", "NewAcc#98765")
    auth.logout()

    # Receptionist also changes their own password
    auth.change_password("reception", "Recep#12345", "NewRecep#98765")
    auth.login("reception", "NewRecep#98765")
    auth.logout()

    # Update access and deactivate user
    auth.login("owner", "Owner#12345")
    auth.update_user_access("acc_user", Role.CUSTOM, permissions=["students", "attendance"], status="inactive")
    with pytest.raises(AuthError, match="موقوف"):
        auth.login("acc_user", "NewAcc#98765")
    with pytest.raises(AuthError, match="موقوف"):
        auth.change_password("acc_user", "NewAcc#98765", "Another#12345")



def test_custom_instructor_receipt_and_statement_pdf(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "INSTRUCTORS_DIR", tmp_path / "instructors_pdf")
    from core.context import AppContext
    from export import service as ex

    session = Session()
    repo = ExcelRepository(tmp_path / "academy.db", tmp_path / "backups", current_user=lambda: session.username)
    repo.ensure_workbook()
    audit = AuditService(repo, lambda: session.username)
    auth = AuthService(repo, audit, session)
    svc = Services(repo, lambda: session.username)
    ctx = AppContext(repo, audit, session, auth, svc=svc)
    auth.first_run_setup(("owner", "Owner#12345"), ("reception", "Recep#12345"))
    auth.login("owner", "Owner#12345")

    iid = svc.instructors.add("د. طارق نبيل", "01011223344", 120)
    g1 = svc.groups.add("بايثون", "برمجة", "مبتدئ", iid, "سبت", "onsite", instructor_hourly_rate=150)
    g2 = svc.groups.add("روبوتكس", "روبوتكس", "متوسط", iid, "أحد", "onsite", instructor_hourly_rate=200)

    s1 = svc.attendance.create_session(g1, TODAY, time(16, 0), time(18, 0))  # 2h * 150 = 300
    s2 = svc.attendance.create_session(g2, TODAY, time(18, 0), time(19, 0))  # 1h * 200 = 200

    # 1. Statement filtered by specific group g1 only
    st_g1 = svc.instructors.statement(iid, TODAY.replace(day=1), TODAY, group_ids=[g1])
    assert st_g1["gross"] == 300 and st_g1["total_hours"] == 2 and "بايثون" in st_g1["group_label"]

    # 2. Statement filtered by specific session s2 + custom item
    cur_m = f"{TODAY:%Y-%m}"
    pdf_path = ex.export_instructor_statement(
        ctx, iid,
        group_ids=[g2],
        months=[cur_m],
        session_ids=[s2],
        custom_title="إيصال صرف مستحقات مدرب",
        custom_item_label="مكافأة تحكيم مسابقة روبوتكس",
        custom_item_amount=500,
        notes="صرف نقدي عن مجموعة روبوتكس + تحكيم المسابقة",
    )
    assert pdf_path.exists() and pdf_path.stat().st_size > 1000

    # 3. Payment receipt PDF
    pid = svc.instructors.record_payment(iid, TODAY.replace(day=1), TODAY, 250, method="cash")
    pay_pdf = ex.export_instructor_statement(ctx, iid, TODAY.replace(day=1), TODAY, payment_id=pid)
    assert pay_pdf.exists() and "pay" in pay_pdf.name


def test_instructor_bonuses_deductions_bilingual_student_names_and_payment_distribution(tmp_path):
    from ui.layout import COLORS
    assert COLORS["primary"] == "#F05500"
    assert COLORS["navy"] == "#0A194E"

    repo = ExcelRepository(tmp_path / "academy.db", tmp_path / "backups", current_user=lambda: "owner")
    repo.ensure_workbook()
    AuditService(repo, lambda: "owner")
    svc = Services(repo, lambda: "owner")

    # 1. Bilingual student name (Arabic + English) + validation
    with pytest.raises(ServiceError, match="بالإنجليزية"):
        svc.students.register(dict(name_ar="عمر أحمد سمير", name_en="Omar", parent_mobile="01011112222"))

    sid1 = svc.students.register(dict(
        name_ar="عمر أحمد سمير",
        name_en="Omar Ahmed Samir",
        parent_mobile="01011112222",
    ))
    st1 = repo.get(Sheet.STUDENTS, student_id=sid1)
    assert st1.name_ar == "عمر أحمد سمير" and st1.name_en == "Omar Ahmed Samir"

    # Search by English name in ledger
    rows_en = svc.ledger().student_rows(query="Omar Ahmed")
    assert len(rows_en) == 1 and rows_en[0]["name_en"] == "Omar Ahmed Samir"

    # 2. Instructor bonuses & deductions
    iid = svc.instructors.add("م. حسام علي", "01055667788", 150)
    gid = svc.groups.add("بايثون", "برمجة", "مبتدئ", iid, "سبت", "onsite", instructor_hourly_rate=150)
    svc.attendance.create_session(gid, TODAY, time(16, 0), time(18, 0))  # 2h * 150 = 300 gross

    lo, hi = TODAY.replace(day=1), TODAY
    svc.instructors.set_adjustments(iid, lo, hi, bonuses=100, deductions=40, notes="مكافأة انتظام وخصم تأخير")
    row = svc.instructors.summary(lo, hi)[0]
    assert row["gross"] == 300 and row["bonuses"] == 100 and row["deductions"] == 40
    assert row["net_due"] == 360 and "مكافأة" in row["notes"]

    st_ins = svc.instructors.statement(iid, lo, hi)
    assert st_ins["bonuses"] == 100 and st_ins["deductions"] == 40 and st_ins["net_due"] == 360

    pid = svc.instructors.record_payment(iid, lo, hi, 360, method="instapay")
    row_after = svc.instructors.summary(lo, hi)[0]
    assert row_after["paid"] == 360 and row_after["net_due"] == 0

    # Reversing payment restores net_due (360) including bonuses & deductions
    svc.instructors.reverse_payment(pid)
    row_rev = svc.instructors.summary(lo, hi)[0]
    assert row_rev["paid"] == 0 and row_rev["bonuses"] == 100 and row_rev["deductions"] == 40 and row_rev["net_due"] == 360

    # 3. Payment method distribution in reports & dashboard today_stats
    eid = svc.students.add_enrollment(sid1, gid, start_date=TODAY, monthly_fee=1000, months=3)
    svc.receipts.create(student_id=sid1, enrollment_id=eid, amount=1000, method="cash", cost_center="بايثون")
    svc.receipts.create(student_id=sid1, enrollment_id=eid, amount=600, method="visa", cost_center="بايثون")
    svc.receipts.create(student_id=sid1, enrollment_id=eid, amount=400, method="instapay", cost_center="بايثون")

    stats = svc.reports.today_stats()
    pm_map = {m["method"]: (m["amount"], m["count"]) for m in stats["payment_methods"]}
    assert pm_map["نقدي"] == (1000.0, 1)
    assert pm_map["فيزا"] == (600.0, 1)
    assert pm_map["إنستا باي"] == (400.0, 1)

    # 4. Daily time period analysis & PDF dashboard/analysis export
    from export.pdf_report import build_report_pdf
    analysis = svc.reports.period_analysis(lo, hi)
    assert analysis["period_from"] == lo and analysis["period_to"] == hi
    assert analysis["totals"]["revenue"] == 2000.0
    assert analysis["totals"]["instructors"] == 300.0
    assert analysis["kpis"]["receipts_count"] == 3
    assert analysis["kpis"]["sessions_held"] == 1
    assert analysis["kpis"]["pending_visa"] == 600.0
    assert len(analysis["by_payment_method"]) == 3
    assert len(analysis["series"]) >= 1 and analysis["series"][0]["revenue"] == 2000.0

    pdf_report = build_report_pdf(tmp_path / "daily_report.pdf", analysis)
    assert pdf_report.exists() and pdf_report.stat().st_size > 1500
    xlsx_report = svc.reports.export_excel(tmp_path / "daily_report.xlsx", lo, hi)
    assert xlsx_report.exists() and xlsx_report.stat().st_size > 1000

    # 5. Admin educational programs management (add & delete programs)
    initial_progs = svc.groups.list_programs()
    assert "بايثون" in initial_progs
    after_add = svc.groups.add_program("تطوير الويب")
    assert "تطوير الويب" in after_add and "تطوير الويب" in svc.groups.list_programs()
    with pytest.raises(ServiceError, match="موجود بالفعل"):
        svc.groups.add_program("تطوير الويب")
    p_stats = {p["name"]: p for p in svc.groups.program_stats()}
    assert p_stats["بايثون"]["groups"] == 1 and p_stats["بايثون"]["students"] == 1
    assert p_stats["تطوير الويب"]["groups"] == 0 and p_stats["تطوير الويب"]["students"] == 0
    after_del = svc.groups.delete_program("تطوير الويب")
    assert "تطوير الويب" not in after_del and "تطوير الويب" not in svc.groups.list_programs()

    # 6. Journal multi-day period view
    svc.journal.add_expense(120, "ضيافة ومشروبات", "ضيافة", on_date=TODAY)
    j_period = svc.journal.period(lo, hi)
    assert j_period["d_from"] == lo and j_period["d_to"] == hi
    assert j_period["day"]["cash_in"] == 1000.0
    assert j_period["day"]["expenses"] == 120.0
    assert j_period["day"]["balance"] == 880.0
    assert len(j_period["cash_receipts"]) == 1 and len(j_period["noncash_receipts"]) == 2

    # 7. Student placement survey & program recommendation engine
    rec_py = svc.students.recommend_program(
        survey={"interest": "coding", "experience": "none", "style": "screen", "computer_skill": "good", "mode": "onsite", "goal": "foundation"},
        age_group="11-14",
    )
    assert rec_py["program"] == "بايثون" and rec_py["track"] == "برمجة" and rec_py["start_level"] == "مبتدئ"
    assert rec_py["suggested_group_id"] == gid

    rec_rob = svc.students.recommend_program(
        survey={"interest": "hardware", "experience": "none", "style": "hands_on", "computer_skill": "basic", "mode": "onsite", "goal": "foundation"},
        age_group="5-7",
    )
    assert rec_rob["program"] == "روبوتيكس" and rec_rob["track"] == "روبوتيكس"

    # 8. Student Graduation Certificate PDF & Student ID Card (QR) PDF
    from export.pdf_certificate import build_certificate_pdf, build_student_id_card_pdf
    cert_pdf = build_certificate_pdf(tmp_path / "cert.pdf", {
        "student_id": sid1, "name_ar": "عمر أحمد سمير", "name_en": "Omar Ahmed Samir",
        "program": "بايثون", "level": "Level 1", "grade": "امتياز", "instructor": "م. حسام علي",
    })
    assert cert_pdf.exists() and cert_pdf.stat().st_size > 1500
    id_pdf = build_student_id_card_pdf(tmp_path / "id_card.pdf", {
        "student_id": sid1, "name_ar": "عمر أحمد سمير", "name_en": "Omar Ahmed Samir",
        "program": "بايثون", "level": "Level 1", "group": "بايثون — Level 1", "parent_mobile": "01011112222",
    })
    assert id_pdf.exists() and id_pdf.stat().st_size > 1500

    # 9. Quick Check-In by Student ID / QR Code
    chk = svc.attendance.quick_check_in(f"FANTASIA-SID:{sid1}", on_date=TODAY)
    assert chk["student_id"] == sid1 and chk["group_id"] == gid and chk["status"] == "present"

    # 10. Automated WhatsApp Message Templates
    from services.common import whatsapp_template_link, whatsapp_template_message
    msg_w = whatsapp_template_message("welcome", student_name="عمر أحمد سمير", program="بايثون", level="Level 1", schedule="السبت 5م")
    assert "عمر أحمد سمير" in msg_w and "بايثون" in msg_w
    wa_url = whatsapp_template_link("01012345678", "receipt", student_name="عمر أحمد سمير", amount="1,000", serial="RC-2026-0001")
    assert wa_url and "wa.me/201012345678" in wa_url

    # 11. Schedule Conflict Detection (Instructor & Student) + Group Level Promotion
    gid_same_sched = svc.groups.add("بايثون", "برمجة", "Level 1", iid, "السبت 5م", "onsite")
    # Update first group schedule so we can test overlap
    svc.groups.update(gid, {"schedule": "السبت والثلاثاء 5م"})
    ins_conflicts = svc.groups.check_schedule_conflicts(iid, "السبت 5م", exclude_group_id=gid_same_sched)
    assert any(c["group_id"] == gid for c in ins_conflicts)
    stu_conflicts = svc.students.check_student_schedule_conflicts(sid1, gid_same_sched)
    assert any(c["group_id"] == gid for c in stu_conflicts)

    promo = svc.groups.promote_group(gid, "Level 2", add_lessons=4, monthly_fee=1200)
    assert promo["new_level"] == "Level 2" and promo["students_promoted"] == 1
    assert svc.ledger().groups[gid].level == "Level 2"

    # 12. Waiting List / Leads CRM
    lead = svc.students.add_lead("ياسين طارق", "01122334455", program="روبوتيكس", notes="يفضل الفترة المسائية")
    assert lead["id"] == 1 and lead["name"] == "ياسين طارق"
    svc.students.update_lead_status(1, "تم التواصل")
    assert svc.students.list_leads()[0]["status"] == "تم التواصل"
    svc.students.delete_lead(1)
    assert len(svc.students.list_leads()) == 0

    # 13. Sibling Detection & Discounted Enrollment
    siblings = svc.students.detect_siblings("01011112222")
    assert any(s.student_id == sid1 for s in siblings)
    sid_sib = svc.students.register(
        {"name_ar": "ليلى أحمد سمير", "parent_mobile": "01011112222"},
        None,
        {"group_id": gid, "start_date": TODAY, "monthly_fee": 1000, "months": 1, "discount_pct": 15},
    )
    sib_enrs = svc.ledger().student_enrollments(sid_sib)
    sib_sum = svc.ledger().summary(sib_enrs[0].enrollment_id)
    assert sib_sum.due == 850.0  # 1000 with 15% sibling discount

    # 14. Monthly Revenue Target & Expense Budget Alerts
    t_res = svc.reports.set_monthly_targets(20000, 100)
    assert t_res["revenue_target"] == 20000.0 and t_res["expense_budget"] == 100.0
    st_dash = svc.reports.today_stats()
    assert st_dash["targets"]["revenue_target"] == 20000.0
    assert st_dash["targets"]["budget_exceeded"] is True

    # 15. One-Click Database Backup Restore
    b_path = repo.backup()
    svc.students.add_lead("مؤقت للحذف بعد الاستعادة", "01200000000")
    assert len(svc.students.list_leads()) == 1
    repo.restore_backup(b_path)
    assert len(svc.students.list_leads()) == 0
