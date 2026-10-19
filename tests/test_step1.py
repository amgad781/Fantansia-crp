"""Step-1 tests:  python -m pytest -q   (from the project root)"""
import os
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

import config
from core.audit import AuditService
from core.auth import AuthError, AuthService, Session
from core.excel_repo import (DataValidationError, DuplicateKeyError, ExcelCorruptError,
                             ExcelFileLockedError, ExcelRepository, RecordNotFoundError)
from core.schemas import (Receipt, Role, Sheet, Student, sheet_columns, validate_entry_date)


@pytest.fixture
def env(tmp_path):
    session = Session()
    repo = ExcelRepository(tmp_path / "academy.xlsx", tmp_path / "backups",
                           current_user=lambda: session.username)
    repo.ensure_workbook()
    audit = AuditService(repo, lambda: session.username)
    auth = AuthService(repo, audit, session)
    return repo, audit, auth, session


def test_workbook_has_all_sheets_and_headers(env):
    repo, *_ = env
    wb = load_workbook(repo.path)
    assert wb.sheetnames == [s.value for s in Sheet] and len(wb.sheetnames) == 17
    for s in Sheet:
        assert [c.value for c in wb[s.value][1]] == sheet_columns(s)
    assert [c.value for c in wb["Receipts"][1]][:4] == ["serial", "date", "student_id", "enrollment_id"]


def test_append_read_roundtrip_and_types(env):
    repo, *_ = env
    repo.append(Sheet.STUDENTS, Student(student_id=1, name_ar="أحمد علي", dob=date(2014, 5, 3),
                                        parent_mobile="٠١٠ ١٢٣٤ ٥٦٧٨", siblings_count=2))
    s = repo.read(Sheet.STUDENTS)[0]
    assert s.dob == date(2014, 5, 3) and s.parent_mobile == "01012345678"
    assert s.status == "active" and s.created_by == "system" and isinstance(s.created_at, datetime)
    assert list(repo.dataframe(Sheet.STUDENTS).columns) == sheet_columns(Sheet.STUDENTS)
    assert repo.dataframe(Sheet.RECEIPTS).empty


def test_duplicate_and_validation(env):
    repo, *_ = env
    repo.append(Sheet.STUDENTS, Student(student_id=1, name_ar="أحمد علي"))
    with pytest.raises(DuplicateKeyError):
        repo.append(Sheet.STUDENTS, Student(student_id=1, name_ar="آخر"))
    with pytest.raises(DataValidationError):
        repo.append(Sheet.STUDENTS, {"student_id": 2, "name_ar": "x"})


def test_receipt_serial_continues_from_220(env):
    repo, *_ = env
    assert repo.next_id(Sheet.RECEIPTS, start=config.RECEIPT_FIRST_SERIAL) == 220
    repo.append(Sheet.RECEIPTS, Receipt(serial=220, date=date.today(), student_id=1, amount=400))
    assert repo.next_id(Sheet.RECEIPTS, start=config.RECEIPT_FIRST_SERIAL) == 221


def test_update_recomputes_derived_and_audits(env):
    repo, *_ = env
    repo.append(Sheet.EMPLOYEE_PAYMENTS, {"id": 1, "employee_id": 1, "month": "2026-10",
                                          "base_salary": 5000, "bonuses": 0, "deductions": 0})
    assert repo.read(Sheet.EMPLOYEE_PAYMENTS)[0].net == 5000
    repo.update(Sheet.EMPLOYEE_PAYMENTS, {"id": 1}, {"bonuses": 500, "deductions": 100})
    assert repo.read(Sheet.EMPLOYEE_PAYMENTS)[0].net == 5400
    log = repo.read(Sheet.AUDIT_LOG)
    assert [a.action for a in log] == ["create", "update"]
    assert '"bonuses": 0.0' in log[1].before and '"bonuses": 500.0' in log[1].after
    with pytest.raises(RecordNotFoundError):
        repo.update(Sheet.EMPLOYEE_PAYMENTS, {"id": 99}, {"bonuses": 1})


def test_session_hours_and_amount(env):
    repo, *_ = env
    repo.append(Sheet.SESSIONS, {"session_id": 1, "date": date(2026, 10, 1), "start_time": "16:00",
                                 "end_time": "17:30", "group_id": 1, "instructor_id": 1, "rate_applied": 100})
    s = repo.read(Sheet.SESSIONS)[0]
    assert s.hours == 1.5 and s.amount == 150 and str(s.start_time) == "16:00:00"
    repo.update(Sheet.SESSIONS, {"session_id": 1}, {"end_time": "18:00"})
    s = repo.read(Sheet.SESSIONS)[0]
    assert s.hours == 2.0 and s.amount == 200


def test_transaction_is_atomic_one_backup(env):
    repo, *_ = env
    n0 = len(repo.list_backups())
    with repo.transaction():
        repo.append(Sheet.STUDENTS, Student(student_id=1, name_ar="أحمد علي"))
        repo.append(Sheet.STUDENTS, Student(student_id=2, name_ar="سارة محمد"))
        assert len(repo.rows(Sheet.STUDENTS)) == 2          # sees uncommitted rows
    assert len(repo.list_backups()) == n0 + 1               # one backup for the whole group
    with pytest.raises(DuplicateKeyError):
        with repo.transaction():
            repo.append(Sheet.STUDENTS, Student(student_id=3, name_ar="مريم أحمد"))
            repo.append(Sheet.STUDENTS, Student(student_id=1, name_ar="مكرر"))
    assert {s.student_id for s in repo.read(Sheet.STUDENTS)} == {1, 2}      # rolled back
    assert not repo.path.with_name("academy.tmp.xlsx").exists()


def test_locked_file_gives_arabic_error_and_no_changes(env, monkeypatch):
    repo, *_ = env
    real_open = open

    def fake_open(file, mode="r", *a, **k):
        if Path(file) == repo.path and "+" in mode:
            raise PermissionError(13, "locked")
        return real_open(file, mode, *a, **k)
    monkeypatch.setattr("builtins.open", fake_open)
    with pytest.raises(ExcelFileLockedError) as ei:
        repo.append(Sheet.STUDENTS, Student(student_id=1, name_ar="أحمد علي"))
    assert "Excel" in str(ei.value) and "أغلق" in str(ei.value)
    monkeypatch.undo()
    assert repo.rows(Sheet.STUDENTS) == []


def test_replace_blocked_gives_locked_error(env, monkeypatch):
    repo, *_ = env
    monkeypatch.setattr("time.sleep", lambda s: None)
    monkeypatch.setattr(os, "replace", lambda a, b: (_ for _ in ()).throw(PermissionError()))
    with pytest.raises(ExcelFileLockedError):
        repo.append(Sheet.STUDENTS, Student(student_id=1, name_ar="أحمد علي"))
    monkeypatch.undo()
    assert not repo.path.with_name("academy.tmp.xlsx").exists()


def test_corrupt_file(env):
    repo, *_ = env
    repo.path.write_bytes(b"not a zip")
    with pytest.raises(ExcelCorruptError):
        repo.rows(Sheet.STUDENTS)


def test_migration_adds_missing_sheet_and_columns(env):
    repo, *_ = env
    wb = load_workbook(repo.path)
    del wb["CashFlow"]
    ws = wb["Students"]
    ws.delete_cols(ws.max_column)                           # drop created_by
    wb.save(repo.path)
    repo.ensure_workbook()
    wb = load_workbook(repo.path)
    assert "CashFlow" in wb.sheetnames
    assert [c.value for c in wb["Students"][1]] == sheet_columns(Sheet.STUDENTS)


def test_backup_pruning(env):
    repo, *_ = env
    repo.backup_dir.mkdir(exist_ok=True)
    now = datetime.now()
    noon = now.replace(hour=12, minute=0, second=0, microsecond=0)
    stamps = {
        "recent_a": now - timedelta(hours=1),
        "recent_b": now - timedelta(hours=2),
        "day3_early": noon - timedelta(days=3, hours=2),     # same day as the next one
        "day3_late": noon - timedelta(days=3) + timedelta(hours=3),
        "too_old": noon - timedelta(days=90),
    }
    for ts in stamps.values():
        (repo.backup_dir / f"academy_{ts:%Y-%m-%d_%H%M%S}_000.xlsx").write_bytes(b"x")
    repo._prune_backups()
    kept = {t.replace(microsecond=0) for t, _ in repo.list_backups()}
    expected = {stamps[k].replace(microsecond=0) for k in ("recent_a", "recent_b", "day3_late")}
    assert kept == expected


def test_auth_flow(env):
    repo, audit, auth, session = env
    assert not auth.has_users()
    with pytest.raises(AuthError):
        auth.first_run_setup(("owner", "123"), ("reception", "reception1"))   # too short
    assert not auth.has_users()
    auth.first_run_setup(("Owner", "owner-pass"), ("reception", "desk-pass"))
    assert {u.username: u.role for u in auth.list_users()} == {"owner": "owner", "reception": "receptionist"}
    with pytest.raises(AuthError):
        auth.first_run_setup(("a1x", "abcdef"), ("b1x", "abcdef"))
    with pytest.raises(AuthError):
        auth.login("owner", "wrong")
    u = auth.login("OWNER", "owner-pass")
    assert u.role == Role.OWNER and session.can("instructors") and session.can("students")
    auth.logout()
    auth.login("reception", "desk-pass")
    assert session.can("students") and not session.can("instructors") and not session.can("employees")
    assert not session.can("dashboard")
    # hashes never reach the audit log
    assert all("$2b$" not in (a.after or "") for a in repo.read(Sheet.AUDIT_LOG))
    actions = [a.action for a in repo.read(Sheet.AUDIT_LOG)]
    assert {"create", "setup", "login_failed", "login", "logout"} <= set(actions)


def test_login_lockout(env):
    repo, audit, auth, session = env
    auth.first_run_setup(("owner", "owner-pass"), ("reception", "desk-pass"))
    for _ in range(config.MAX_LOGIN_ATTEMPTS):
        with pytest.raises(AuthError):
            auth.login("owner", "bad")
    with pytest.raises(AuthError) as ei:
        auth.login("owner", "owner-pass")                   # correct, but locked
    assert "محاولات" in str(ei.value)


def test_change_password(env):
    repo, audit, auth, session = env
    auth.first_run_setup(("owner", "owner-pass"), ("reception", "desk-pass"))
    auth.change_password("owner", "owner-pass", "new-secret")
    auth.login("owner", "new-secret")


def test_entry_date_validation():
    today = date(2026, 10, 12)
    assert validate_entry_date(date(2026, 10, 1), today=today)
    for bad in (date(2020, 2, 1), date(2026, 10, 13)):
        with pytest.raises(ValueError):
            validate_entry_date(bad, today=today)
