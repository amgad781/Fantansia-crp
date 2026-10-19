"""Business-logic tests (services layer):  python -m pytest -q tests/test_services.py"""
from datetime import date, time, timedelta

import pytest

import config
from core.audit import AuditService
from core.auth import AuthService, Session
from core.excel_repo import ExcelRepository
from core.schemas import Sheet
from services.common import ServiceError, month_key, normalize_ar, shift_month, whatsapp_link
from services.registry import Services

TODAY = date.today()
M0 = month_key(TODAY)


@pytest.fixture
def svc(tmp_path):
    session = Session()
    repo = ExcelRepository(tmp_path / "academy.xlsx", tmp_path / "backups", current_user=lambda: "reception")
    repo.ensure_workbook()
    AuditService(repo, lambda: "reception")
    return Services(repo, lambda: "reception")


@pytest.fixture
def world(svc):
    """1 instructor (120/h), 1 group, 1 student enrolled with 2 monthly fees of 1500."""
    iid = svc.instructors.add("كريم حسن", "01199998888", 120)
    gid = svc.groups.add("بايثون أونلاين", "برمجة", "مبتدئ", iid, "سبت وثلاثاء", "online")
    sid = svc.students.register(
        dict(name_ar="عمر أحمد سمير", parent_mobile="01012345678", age_group="8-10"),
        assessment=dict(technical_grade="4", track="برمجة", start_level="مبتدئ"),
        enrollment=dict(group_id=gid, start_date=TODAY, monthly_fee=1500, months=2))
    eid = svc.ledger().student_enrollments(sid)[0].enrollment_id
    return svc, iid, gid, sid, eid


def eid_of(svc, sid):
    return svc.ledger().student_enrollments(sid)[0].enrollment_id


def test_helpers():
    assert normalize_ar("أَحمد إبراهيم آل") == "احمد ابراهيم ال"
    assert shift_month("2026-11", 3) == "2027-02"
    assert whatsapp_link("010 1234 5678").startswith("https://wa.me/201012345678")
    assert whatsapp_link("123") is None


def test_register_creates_everything(world):
    svc, iid, gid, sid, eid = world
    assert len(svc.repo.read(Sheet.FEE_SCHEDULE)) == 2
    assert svc.students.latest_assessment(sid).track == "برمجة"
    with pytest.raises(ServiceError):                              # same name + mobile
        svc.students.register(dict(name_ar="عمر أحمد سمير", parent_mobile="01012345678"))
    with pytest.raises(ServiceError):                              # one-word name
        svc.students.register(dict(name_ar="عمر"))
    with pytest.raises(ServiceError):                              # duplicate enrollment
        svc.students.add_enrollment(sid, gid, TODAY, 1500)
    with pytest.raises(ServiceError, match="مسجل بالفعل"):         # duplicate fee line: friendly message
        svc.students.add_fee(eid_of(svc, sid), "monthly", M0, 100)
    svc.students.add_fee(eid_of(svc, sid), "other", M0, 100)       # different item is fine


def test_receipt_numbering_lessons_and_fifo(world):
    svc, iid, gid, sid, eid = world
    r1 = svc.receipts.create(student_id=sid, enrollment_id=eid, amount=1000, cost_center="بايثون أونلاين")
    assert r1.serial == config.RECEIPT_FIRST_SERIAL                # continues from the paper pad
    led = svc.ledger()
    sm = led.summary(eid)
    assert sm.paid == 1000 and sm.remaining_fees == 2000 and sm.purchased == 0   # month 1 not fully paid
    r2 = svc.receipts.create(student_id=sid, enrollment_id=eid, amount=500)
    sm = svc.ledger().summary(eid)
    assert r2.serial == 221 and sm.lines[0].state == "paid" and sm.purchased == config.LESSONS_PER_PAID_MONTH
    assert svc.receipts.remaining_after(220) == 2000 and svc.receipts.remaining_after(221) == 1500
    # cancel the second receipt -> month 1 is no longer fully paid -> lessons go back
    with pytest.raises(ServiceError):
        svc.receipts.reverse(221, "")
    svc.receipts.reverse(221, "خطأ في المبلغ")
    sm = svc.ledger().summary(eid)
    assert sm.paid == 1000 and sm.purchased == 0
    with pytest.raises(ServiceError):
        svc.receipts.reverse(221, "تاني")
    assert [r for r in svc.repo.read(Sheet.RECEIPTS) if r.serial == 221][0].status == "reversed"


def test_receipt_validation(world):
    svc, iid, gid, sid, eid = world
    for bad in (0, -5, None):
        with pytest.raises(ServiceError):
            svc.receipts.create(student_id=sid, enrollment_id=eid, amount=bad)
    with pytest.raises(ServiceError):                              # enrollment required
        svc.receipts.create(student_id=sid, amount=100)
    with pytest.raises(ServiceError):                              # typo date
        svc.receipts.create(student_id=sid, enrollment_id=eid, amount=100, on_date=date(2020, 1, 2))
    with pytest.raises(ServiceError):
        svc.receipts.create(student_id=sid, enrollment_id=eid, amount=100, on_date=TODAY + timedelta(days=3))


def test_sessions_attendance_remaining_lessons_and_absence_limit(world):
    svc, iid, gid, sid, eid = world
    svc.repo.update(Sheet.ENROLLMENTS, {"enrollment_id": eid}, {"start_date": TODAY - timedelta(days=10)})
    svc.receipts.create(student_id=sid, enrollment_id=eid, amount=1500)          # 4 lessons
    s1 = svc.attendance.create_session(gid, TODAY, time(17, 0), time(18, 30))
    with pytest.raises(ServiceError):                                             # overlapping
        svc.attendance.create_session(gid, TODAY, time(18, 0), time(19, 0))
    s2 = svc.attendance.create_session(gid, TODAY - timedelta(days=1), time(17, 0), time(18, 0))
    s3 = svc.attendance.create_session(gid, TODAY - timedelta(days=2), time(17, 0), time(18, 0))
    svc.attendance.mark(s1, sid, "present")
    svc.attendance.mark(s2, sid, "absent")
    svc.attendance.mark(s3, sid, "excused")                                       # does not consume
    sm = svc.ledger().summary(eid)
    assert (sm.purchased, sm.consumed, sm.remaining_lessons, sm.absences) == (4, 2, 2, 1)
    assert not sm.absence_exceeded
    s4 = svc.attendance.create_session(gid, TODAY - timedelta(days=3), time(17, 0), time(18, 0))
    svc.attendance.mark(s4, sid, "absent")
    assert svc.ledger().summary(eid).absence_exceeded                             # 2 absences > limit 1
    # sessions before the enrollment start never count
    s0 = svc.attendance.create_session(gid, TODAY - timedelta(days=30), time(9, 0), time(10, 0))
    svc.attendance.mark(s0, sid, "present")
    assert svc.ledger().summary(eid).consumed == 3        # s1 present, s2 absent, s4 absent; s0 ignored
    # a cancelled session consumes nothing
    svc.attendance.set_session_status(s4, "cancelled")
    sm = svc.ledger().summary(eid)
    assert sm.consumed == 2 and sm.absences == 1
    assert svc.attendance.next_status(None) == "present" and svc.attendance.next_status("excused") == "present"
    g = svc.attendance.grid(gid, M0)
    assert len(g["sessions"]) >= 1 and g["students"][0]["student_id"] == sid


def test_student_rows_search_and_badges(world):
    svc, iid, gid, sid, eid = world
    svc.receipts.create(student_id=sid, enrollment_id=eid, amount=1500)
    rows = svc.ledger().student_rows()
    assert len(rows) == 1 and rows[0]["lessons"] == 4 and rows[0]["lessons_color"] == "positive"
    assert rows[0]["paid"] == 1500 and rows[0]["balance"] == 1500
    L = svc.ledger
    assert L().student_rows(query="عمر احمد") and L().student_rows(query="0101234")
    assert not L().student_rows(query="زيد")
    assert not L().student_rows(low_only=True)
    assert L().student_rows(program="بايثون أونلاين") and not L().student_rows(program="أردوينو")


def test_instructor_hourly_salary_to_date(world):
    svc, iid, gid, sid, eid = world
    first = TODAY.replace(day=1)
    svc.attendance.create_session(gid, TODAY - timedelta(days=1), time(17, 0), time(18, 30))   # 1.5h * 120
    svc.attendance.create_session(gid, TODAY - timedelta(days=2), time(10, 0), time(11, 0))    # 1h * 120
    fut = svc.attendance.create_session(gid, TODAY + timedelta(days=5), time(17, 0), time(18, 0))
    svc.instructors.change_rate(iid, 150)                      # later raise must not touch frozen sessions
    last = date(TODAY.year + (TODAY.month == 12), TODAY.month % 12 + 1, 1) - timedelta(days=1)
    lo = min(first, TODAY - timedelta(days=2))
    row = svc.instructors.summary(lo, last)[0]
    assert row["hours"] == 2.5 and row["gross"] == 300 and row["rate"] == 150   # future session excluded
    assert svc.instructors.rate_for(iid, TODAY - timedelta(days=1)) == 120
    with pytest.raises(ServiceError):
        svc.instructors.record_payment(iid, lo, last, 301)    # more than due
    pid = svc.instructors.record_payment(iid, lo, last, 100)
    row = svc.instructors.summary(lo, last)[0]
    assert row["paid"] == 100 and row["net_due"] == 200
    st = svc.instructors.statement(iid, lo, last)
    assert st["gross"] == 300 and st["previous_paid"] == 100 and st["net_due"] == 200 and len(st["lines"]) == 2
    svc.instructors.reverse_payment(pid)
    assert svc.instructors.summary(lo, last)[0]["net_due"] == 300
    # cancelling a session removes its pay
    s = svc.repo.read(Sheet.SESSIONS)[0]
    svc.attendance.set_session_status(s.session_id, "postponed")
    assert svc.instructors.summary(lo, last)[0]["gross"] < 300


def test_collection_overdue_and_whatsapp(world):
    svc, iid, gid, sid, eid = world
    sid2 = svc.students.register(dict(name_ar="سلمى محمود علي", parent_whatsapp="01155556666"),
                                 enrollment=dict(group_id=gid, start_date=TODAY - timedelta(days=65),
                                                 monthly_fee=1000, months=3))
    od = svc.ledger().overdue_list()
    names = {o["name"]: o for o in od}
    assert "سلمى محمود علي" in names and names["سلمى محمود علي"]["amount"] >= 2000   # 2 past/current months
    assert whatsapp_link(names["سلمى محمود علي"]["mobile"]).startswith("https://wa.me/2011555")
    months = [M0, shift_month(M0, 1)]
    grid = svc.ledger().collection_grid(months)
    assert len(grid) == 2 and grid[0]["cells"][M0]["state"] in ("unpaid", "partial", "paid", "none")


def test_employees_month_flow(svc):
    eid = svc.employees.add("منى السيد", "موظفة استقبال", "01000000000", 4000)
    rows = svc.employees.month_rows(M0)
    assert rows[0]["status"] == "unsaved" and rows[0]["net"] == 4000
    svc.employees.set_adjustments(eid, M0, 500, 200)
    r = svc.employees.month_rows(M0)[0]
    assert r["net"] == 4300 and r["status"] == "pending"
    svc.employees.pay(eid, M0, "instapay")
    r = svc.employees.month_rows(M0)[0]
    assert r["status"] == "paid" and r["method"] == "instapay"
    with pytest.raises(ServiceError):
        svc.employees.pay(eid, M0)
    with pytest.raises(ServiceError):
        svc.employees.set_adjustments(eid, M0, 0, 0)
    svc.employees.reverse(r["payment_id"])
    assert svc.employees.payslip_data(eid, M0)["net"] == 4000          # fresh pending row snapshot

    # Direct bonuses and deductions via pay()
    eid2 = svc.employees.add("أحمد علي", "أمن", "01100000000", 3000)
    svc.employees.pay(eid2, M0, "cash", bonuses=400, deductions=100)
    r2 = [x for x in svc.employees.month_rows(M0) if x["employee_id"] == eid2][0]
    assert r2["status"] == "paid" and r2["bonuses"] == 400 and r2["deductions"] == 100 and r2["net"] == 3300


def test_journal_drawer_and_reports(world):
    svc, iid, gid, sid, eid = world
    svc.receipts.create(student_id=sid, enrollment_id=eid, amount=1500, method="cash", cost_center="بايثون أونلاين")
    svc.receipts.create(student_id=sid, enrollment_id=eid, amount=500, method="instapay", cost_center="بايثون أونلاين")
    svc.journal.add_expense(200, "أدوات مكتبية", "أدوات ومستلزمات")
    svc.journal.add_cashflow("petty_cash", 100)
    svc.journal.add_cashflow("handover", 700, receipt_no=220)
    day = svc.journal.day(TODAY)
    assert day["receipts_total"] == 2000 and len(day["noncash_receipts"]) == 1
    assert day["day"]["balance"] == 1500 - 200 - 100 - 700 == 500 and svc.journal.drawer_balance() == 500
    xid = svc.repo.read(Sheet.EXPENSES)[0].id
    svc.journal.reverse_expense(xid)
    assert svc.journal.drawer_balance() == 700
    with pytest.raises(ServiceError):
        svc.journal.add_expense(0, "x")
    with pytest.raises(ServiceError):
        svc.journal.add_expense(50, "", "أخرى")

    ms = svc.reports.month_summary(M0)
    assert ms["revenue"] == 2000 and ms["expenses"] == 0
    bp = svc.reports.by_program(M0, M0)[0]
    assert bp["program"] == "بايثون أونلاين" and bp["revenue"] == 2000
    assert len(svc.reports.series(M0, 6)) == 6
    stats = svc.reports.today_stats()
    assert stats["revenue_today"] == 2000 and stats["drawer"] == 700
    out = svc.reports.export_excel(svc.repo.path.parent / "rep.xlsx", M0, M0)
    from openpyxl import load_workbook
    assert "ملخص" in load_workbook(out).sheetnames


def test_journal_expense_categories(world):
    svc, iid, gid, sid, eid = world
    cats = svc.journal.list_categories()
    assert "أخرى" in cats
    assert "إيجار" in cats

    # Add a new category
    updated = svc.journal.add_category("بوفيه وضيافة خاصة")
    assert "بوفيه وضيافة خاصة" in updated
    assert updated[-1] == "أخرى"

    # Persistence in repo settings
    assert "بوفيه وضيافة خاصة" in svc.journal.list_categories()

    # Reject duplicate (even with slight Arabic variations)
    with pytest.raises(ServiceError):
        svc.journal.add_category("بوفيه وضيافه خاصه")

    # Reject too short
    with pytest.raises(ServiceError):
        svc.journal.add_category(" أ ")

    # Add expense using the new category
    xid = svc.journal.add_expense(150, "شاي وسكر", "بوفيه وضيافة خاصة")
    assert xid > 0
    exp = svc.repo.get(Sheet.EXPENSES, id=xid)
    assert exp.category == "بوفيه وضيافة خاصة"

    # Cannot delete "أخرى"
    with pytest.raises(ServiceError):
        svc.journal.delete_category("أخرى")

    # Cannot delete category in active use
    with pytest.raises(ServiceError):
        svc.journal.delete_category("بوفيه وضيافة خاصة")

    # Add temporary category and delete it
    svc.journal.add_category("تصنيف مؤقت")
    assert "تصنيف مؤقت" in svc.journal.list_categories()
    svc.journal.delete_category("تصنيف مؤقت")
    assert "تصنيف مؤقت" not in svc.journal.list_categories()
